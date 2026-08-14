# -*- coding: utf-8 -*-
"""Rebuild web/data/places.js from the UPDATED workbook (112 places, 2 sheets).

Sheet 1 "Бангкок 11-15 окт" — place details (new numbering).
Sheet 2 "По районам"        — the same places grouped by district.

Coordinates are carried over from the previous dataset by English name and
extended with curated coordinates for the newly added places.
"""
import json, re, os, sys

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = "/root/.claude/uploads/21e2b459-7194-5ed7-9d02-84c3d141cbe9/3f929085-______________1115_____1.xlsx"
OLD = os.path.join(HERE, "..", "web", "data", "places.js")
OUT = os.path.join(HERE, "..", "web", "data", "places.js")

import openpyxl

# ---- reuse the parsing helpers from the first builder -----------------------
sys.path.insert(0, HERE)
from build_data import (parse_hours, parse_trip_dates, parse_price, parse_flags,
                        CAT_KEY, CAT_DAYPART, THAI)

# Curated coordinates for places that are new in this workbook (by EN name).
NEW_COORDS = {
    "Topkart Bangkok": [100.5560, 13.7515],
    "The Ancient City (Muang Boran)": [100.6060, 13.5330],
    "Dusit Central Park (Central Park mall)": [100.5350, 13.7290],
    "1981 Soul&Sold": [100.6220, 13.7570],
    "Train Night Market Srinagarindra": [100.6470, 13.6960],
    "Ozone One Market": [100.5930, 13.9200],
    "Talat Noi": [100.5120, 13.7345],
    "Mega Plaza Saphan Lek": [100.5010, 13.7455],
    "Kodtalay The Riverfront Seafood Buffet": [100.5060, 13.6990],
    "Check In Bar (Soi 10)": [100.5540, 13.7420],
    "Golden Dome Cabaret Show": [100.5720, 13.7690],
    "Calypso Cabaret": [100.5027, 13.7043],
    "STUSSY Bangkok Chapter Store": [100.5470, 13.7440],
    "2nd STREET CentralWorld": [100.5396, 13.7466],
    "Rajamangala National Stadium": [100.6220, 13.7555],
}

def load_old_coords():
    """name_en -> [lon,lat] from the previously generated dataset."""
    try:
        txt = open(OLD, encoding="utf-8").read()
        data = json.loads(txt.split("window.PLACES = ", 1)[1].rsplit(";", 1)[0])
        return {p["nameEn"]: [p["lon"], p["lat"]] for p in data}
    except Exception:
        return {}

def districts_map(wb):
    """Read sheet 2: returns id -> (district, subgroup) and district order."""
    ws = wb["По районам"]
    out, order, cur, sub = {}, [], None, None
    for row in ws.iter_rows(values_only=True):
        first = (str(row[0]).strip() if row[0] is not None else "")
        if not first:
            continue
        if first.isdigit():
            out[int(first)] = (cur, sub)
            continue
        # heading rows
        clean = first.replace("▍", "").strip()
        if not clean or clean.startswith("Те же") or clean.startswith("№"):
            continue
        if clean.startswith("Можно посетить также"):
            sub = clean.replace("Можно посетить также —", "").strip(" —")
            continue
        if first.startswith("      "):   # indented sub-heading
            sub = clean
            continue
        if clean.startswith("В вашем списке"):
            continue
        cur, sub = clean, None
        if cur not in order:
            order.append(cur)
    return out, order

wb = openpyxl.load_workbook(XLSX, data_only=True)
old = load_old_coords()
dmap, dorder = districts_map(wb)

ws = wb["Бангкок 11-15 окт"]
cat = None
places, missing = [], []
for row in ws.iter_rows(values_only=True):
    first = (str(row[0]).strip() if row[0] is not None else "")
    if not first:
        continue
    if first.startswith("▍"):
        cat = first.replace("▍", "").strip()
        continue
    if not first.isdigit():
        continue
    i = int(first)
    name_en = (row[1] or "").strip()
    lonlat = old.get(name_en) or NEW_COORDS.get(name_en)
    if not lonlat:
        missing.append((i, name_en))
        continue
    catk = CAT_KEY.get(cat, "sight")
    price = parse_price(row[5] or "")
    district, sub = dmap.get(i, (None, None))
    places.append({
        "id": i, "cat": catk, "catName": cat,
        "district": district, "districtSub": sub,
        "nameEn": name_en, "nameRu": (row[2] or "").strip(),
        "nameTh": THAI.get(0, ""),  # filled below by name match
        "address": (row[3] or "").strip(), "desc": (row[4] or "").strip(),
        "priceText": (row[5] or "").strip(), "price": price,
        "hoursText": (row[6] or "").strip(), "hours": parse_hours(row[6] or "", catk),
        "features": (row[7] or "").strip(),
        "availabilityText": (row[8] or "").strip(),
        "trip": parse_trip_dates(row[8] or "", row[6] or ""),
        "daypart": CAT_DAYPART.get(catk, "any"),
        "flags": parse_flags(row[7] or "", row[8] or "", row[5] or "", name_en),
        "lon": lonlat[0], "lat": lonlat[1],
    })

# Thai names: carry over from the old dataset by English name
try:
    oldtxt = open(OLD, encoding="utf-8").read()
    olddata = json.loads(oldtxt.split("window.PLACES = ", 1)[1].rsplit(";", 1)[0])
    th = {p["nameEn"]: p.get("nameTh", "") for p in olddata if p.get("nameTh")}
except Exception:
    th = {}
for p in places:
    p["nameTh"] = th.get(p["nameEn"], "")

print("parsed:", len(places), "places")
print("districts:", dorder)
if missing:
    print("MISSING COORDS:")
    for m in missing:
        print("   ", m)

if not missing:
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("// Auto-generated by scripts/build_data2.py — Bangkok places (112).\n")
        f.write("window.PLACES = ")
        json.dump(places, f, ensure_ascii=False, indent=0)
        f.write(";\n")
    print("wrote", OUT)
